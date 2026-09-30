"""File Excel xuất ra (spec ``chat-report-export``, b10 task 4.3, 4.4, 4.7)."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dms.chat.ai.export.export_store import ExportStore, safe_username
from dms.chat.ai.export.xlsx_writer import (
    AI_COMMENTARY_NOTE,
    ExportInfo,
    ExportTooLarge,
    sanitize_cell,
    sheet_title,
    write_workbook,
)

NOW = datetime(2026, 9, 15, 10, 0, tzinfo=UTC)


def info(**overrides) -> ExportInfo:
    data = {
        "title": "Báo cáo tuần",
        "exported_by": "Nguyễn Văn A",
        "exported_at": "15/09/2026 10:00",
        "scope": "Truyền thống Vùng 1",
        "date_range": "07/09/2026 – 13/09/2026",
    }
    data.update(overrides)
    return ExportInfo(**data)


def ranking_block(items: list[dict], title: str = "Xếp hạng sản phẩm") -> dict:
    return {
        "kind": "ranking",
        "title": title,
        "subtitle": "07/09/2026 – 13/09/2026",
        "payload": {
            "items": items,
            "label_title": "Sản phẩm",
            "value_title": "Số vấn đề",
        },
    }


def quote_block(quotes: list[dict]) -> dict:
    return {"kind": "quote", "title": "Phản hồi tiêu biểu", "payload": {"quotes": quotes}}


# ── Kiểu ô ──


def test_numbers_and_percentages_are_numeric_cells(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    write_workbook(
        path, info(), [ranking_block([{"rank": 1, "label": "Đèn", "value": 1234, "percent": 12.5}])]
    )
    sheet = load_workbook(path)["Xếp hạng sản phẩm"]
    count_cell = sheet.cell(row=5, column=3)
    percent_cell = sheet.cell(row=5, column=4)
    assert count_cell.value == 1234 and count_cell.data_type == "n"
    assert count_cell.number_format == "#,##0"
    assert percent_cell.value == pytest.approx(0.125)
    assert percent_cell.number_format == "0.0%"


def test_dates_are_date_cells(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    write_workbook(
        path,
        info(),
        [quote_block([{"issue_code": "NT-1", "content": "x", "issue_date": "2026-09-09"}])],
    )
    cell = load_workbook(path)["Phản hồi tiêu biểu"].cell(row=5, column=4)
    assert cell.value == datetime(2026, 9, 9)
    assert cell.number_format == "dd/mm/yyyy"


def test_missing_values_stay_text(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    write_workbook(path, info(), [ranking_block([{"rank": 1, "label": "Đèn", "value": None}])])
    cell = load_workbook(path)["Xếp hạng sản phẩm"].cell(row=5, column=3)
    assert cell.value == "Chưa đủ dữ liệu" and cell.data_type == "s"


# ── Chặn công thức ──


@pytest.mark.parametrize(
    "hostile",
    ['=HYPERLINK("http://x","bấm")', "+cmd|' /C calc", "-1+1", "@SUM(A1)", "\tx", "\rx"],
)
def test_formula_like_text_is_escaped(hostile: str):
    assert sanitize_cell(hostile).startswith("'")


def test_workbook_has_no_formula_cells(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    write_workbook(
        path,
        info(notes=("=1+1",), commentary=("@risky",)),
        [
            ranking_block([{"rank": 1, "label": '=HYPERLINK("http://x","bấm")', "value": 5}]),
            quote_block([{"issue_code": "NT-1", "content": "+cmd|' /C calc"}]),
        ],
    )
    workbook = load_workbook(path)
    assert not any(
        cell.data_type == "f" for sheet in workbook for row in sheet.iter_rows() for cell in row
    )
    assert load_workbook(path)["Xếp hạng sản phẩm"].cell(row=5, column=2).value.startswith("'=")


# ── Sheet ──


def test_info_sheet_is_first_and_lists_context(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    write_workbook(
        path,
        info(
            assumptions=("Hiểu là tuần trước.",),
            highlights=("Số vấn đề tăng 23%.",),
            commentary=("Đèn LED Bulb dẫn đầu.",),
        ),
        [ranking_block([{"rank": 1, "label": "Đèn", "value": 5}])],
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames[0] == "Thông tin"
    text = "\n".join(
        str(cell.value or "") for row in workbook["Thông tin"].iter_rows() for cell in row
    )
    for expected in (
        "Nguyễn Văn A",
        "15/09/2026 10:00",
        "Truyền thống Vùng 1",
        "Hiểu là tuần trước.",
        "Số vấn đề tăng 23%.",
        "Đèn LED Bulb dẫn đầu.",
        AI_COMMENTARY_NOTE,
    ):
        assert expected in text


def test_long_and_duplicate_sheet_titles():
    used: set[str] = set()
    long_title = "Phân bổ theo sản phẩm của toàn bộ đơn vị trong kỳ"
    first = sheet_title(long_title, used)
    second = sheet_title(long_title, used)
    assert first != second
    assert len(first) <= 31 and len(second) <= 31


def test_invalid_sheet_characters_are_removed():
    assert "/" not in sheet_title("Sản phẩm / Đơn vị: [2026]", set())


# ── Giới hạn ──


def test_rows_are_capped_and_marked_truncated(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    items = [{"rank": i, "label": f"SP {i}", "value": i} for i in range(1, 7201)]
    result = write_workbook(path, info(), [ranking_block(items)], max_rows=5000)
    assert result.rows_exported == 5000 and result.truncated is True
    sheet = load_workbook(path)["Xếp hạng sản phẩm"]
    assert sheet.max_row == 5004  # 1 tiêu đề + 1 phụ đề + 1 trống + 1 header + 5000 dòng
    text = "\n".join(
        str(cell.value or "")
        for row in load_workbook(path)["Thông tin"].iter_rows()
        for cell in row
    )
    assert "5.000" in text


def test_oversized_file_is_removed(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    items = [{"rank": i, "label": f"SP {i}", "value": i} for i in range(1, 500)]
    with pytest.raises(ExportTooLarge):
        write_workbook(path, info(), [ranking_block(items)], max_bytes=1024)
    assert not path.exists()
    assert not list(tmp_path.glob("*.part"))


def test_write_is_atomic(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    write_workbook(path, info(), [ranking_block([{"rank": 1, "label": "Đèn", "value": 5}])])
    assert path.is_file() and not list(tmp_path.glob("*.part"))


# ── ExportStore ──


def test_safe_username_only_uses_allowed_characters():
    folded = safe_username("Nguyễn Văn A/../root")
    assert all(char.isalnum() or char in "-_" for char in folded)
    assert safe_username("a") != safe_username("A")  # hash giữ cho hai user khác nhau


def store(tmp_path: Path, **kwargs) -> ExportStore:
    return ExportStore(tmp_path / "chat_exports", **kwargs)


def make_export(st: ExportStore, owner: str, *, now: datetime = NOW) -> str:
    export_id, path = st.reserve(owner)
    path.write_bytes(b"xlsx")
    st.save_metadata(owner, export_id, filename="a.xlsx", size_bytes=4, now=now)
    return export_id


def test_owner_can_read_but_other_user_cannot(tmp_path: Path):
    st = store(tmp_path)
    export_id = make_export(st, "an")
    assert st.get("an", export_id, now=NOW) is not None
    assert st.get("binh", export_id, now=NOW) is None


def test_expired_export_is_not_returned(tmp_path: Path):
    st = store(tmp_path, ttl_hours=1)
    export_id = make_export(st, "an")
    assert st.get("an", export_id, now=NOW + timedelta(hours=2)) is None


@pytest.mark.parametrize("bad", ["../settings", "zz", "", "a" * 31, "A" * 32])
def test_invalid_export_ids_are_rejected(tmp_path: Path, bad: str):
    assert store(tmp_path).get("an", bad, now=NOW) is None


def test_cleanup_removes_both_files(tmp_path: Path):
    st = store(tmp_path, ttl_hours=1)
    export_id = make_export(st, "an")
    assert st.cleanup(NOW + timedelta(hours=2)) == 1
    xlsx, meta = st.paths_for("an", export_id)
    assert not xlsx.exists() and not meta.exists()


def test_cleanup_keeps_fresh_files(tmp_path: Path):
    st = store(tmp_path, ttl_hours=24)
    make_export(st, "an")
    assert st.cleanup(NOW + timedelta(hours=1)) == 0


def test_limit_removes_the_oldest_export(tmp_path: Path):
    st = store(tmp_path, max_active_per_user=3)
    ids = [make_export(st, "an", now=NOW + timedelta(minutes=i)) for i in range(4)]
    remaining = {record.export_id for record in st.records_for("an")}
    assert len(remaining) == 3
    assert ids[0] not in remaining and ids[-1] in remaining
