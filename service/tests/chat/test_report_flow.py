"""Luồng báo cáo đầu-cuối (spec ``chat-report-composition``, b10 task 3.5).

Planner giả chỉ trả intent; bước do template dựng, số liệu do ``ReportFakeExecutor`` trả.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime

import pytest

from dms.chat.ai.answer_events import DoneStatus, EventType, ListSink
from dms.chat.ai.orchestrator import ChatOrchestrator, OrchestratorConfig
from dms.chat.ai.planner_config import PlannerConfig
from dms.chat.ai.query_planner import QueryPlanner
from dms.chat.ai.response_shaper import AnswerComposer, ComposerConfig
from dms.chat.ai.schema_retriever import SchemaRetriever
from dms.chat.ai.types import (
    Decision,
    Reason,
    SectionStatus,
    TurnMode,
    TurnRequest,
)
from dms.chat.ai.understanding import UnderstandingConfig
from dms.chat.contract import UserScope
from dms.chat.guardrails.plan_guard import PlanGuard, PlanGuardConfig

from .ai_fakes import FixedClock, ScriptedLLM, StaticMetadataProvider
from .report_fakes import TV1, TV2, ReportFakeExecutor

CLOCK = FixedClock(datetime(2026, 9, 15, 3, 0, tzinfo=UTC))
M5_PATTERNS = frozenset({"sql_template", "fts5_search", "semantic_view"})


def plan_json(intent: str, sections: list[str] | None = None) -> str:
    return json.dumps(
        {
            "intent": intent,
            "answer_shape": "table",
            "confidence": 0.95,
            "steps": [],
            "report_sections": sections or [],
        },
        ensure_ascii=False,
    )


def build(
    executor,
    *,
    plan: str,
    config: OrchestratorConfig | None = None,
    stream_chunks=None,
) -> tuple[ChatOrchestrator, ScriptedLLM]:
    llm = ScriptedLLM(default=plan, stream_chunks=stream_chunks)
    metadata = StaticMetadataProvider()
    planner_config = PlannerConfig(milestone="M5", enabled_patterns=M5_PATTERNS)
    return (
        ChatOrchestrator(
            llm=llm,
            planner=QueryPlanner(llm, SchemaRetriever(metadata, config=planner_config)),
            plan_guard=PlanGuard(
                metadata,
                config=PlanGuardConfig(milestone="M5", enabled_patterns=M5_PATTERNS),
                clock=CLOCK,
            ),
            executor=executor,
            metadata=metadata,
            config=config or OrchestratorConfig(),
            understanding=UnderstandingConfig(),
            clock=CLOCK,
        ),
        llm,
    )


def turn(question: str, *units: str, admin: bool = False, **kwargs) -> TurnRequest:
    scope = (
        UserScope(username="admin", role="admin", display_name="Admin", unit_ids=[])
        if admin
        else UserScope(username="nv", role="user", display_name="NV", unit_ids=list(units))
    )
    return TurnRequest(question=question, scope=scope, session_id="s1", **kwargs)


def compose(outcome, *, llm=None, commentary: bool = False) -> ListSink:
    sink = ListSink()
    composer = AnswerComposer(
        llm=llm,
        config=ComposerConfig(commentary_enabled=commentary, milestone="M5"),
    )
    composer.compose(outcome, sink)
    return sink


def blocks(sink: ListSink) -> list[dict]:
    return [event.data for event in sink.of_type(EventType.DATA_BLOCK)]


# ── Orchestrator ──


def test_weekly_report_runs_every_section():
    executor = ReportFakeExecutor()
    orchestrator, _ = build(executor, plan=plan_json("REPORT_WEEKLY"))

    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))

    assert outcome.decision is Decision.RUN
    assert outcome.mode is TurnMode.REPORT
    assert [section.section_id for section in outcome.sections] == [
        "overview",
        "trend",
        "products",
        "geography",
        "units",
        "priority",
    ]
    assert all(section.status is SectionStatus.OK for section in outcome.sections)
    assert outcome.report_range is not None
    assert outcome.report_range.date_from.isoformat() == "2026-09-07"


def test_default_weekly_range_records_assumption():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_WEEKLY"))
    outcome = orchestrator.handle(turn("Báo cáo tuần", TV1))
    assert outcome.report_range is not None
    assert outcome.report_range.date_from.isoformat() == "2026-09-07"
    assert any("tuần trước" in text for text in outcome.assumptions)


def test_failed_section_does_not_stop_the_report():
    executor = ReportFakeExecutor(failures=["get_geography"])
    orchestrator, _ = build(executor, plan=plan_json("REPORT_WEEKLY"))

    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    by_id = {section.section_id: section.status for section in outcome.sections}

    assert by_id["geography"] is SectionStatus.ERROR
    assert by_id["products"] is SectionStatus.OK and by_id["priority"] is SectionStatus.OK
    sink = compose(outcome)
    assert sink.done is not None
    assert sink.done.data["status"] == DoneStatus.PARTIAL.value
    failed = [b for b in blocks(sink) if b.get("section", {}).get("id") == "geography"]
    assert failed and failed[0]["subtitle"] == "Không lấy được dữ liệu"


def test_failed_overview_fails_the_turn():
    executor = ReportFakeExecutor(failures=["get_overview"])
    orchestrator, _ = build(executor, plan=plan_json("REPORT_WEEKLY"))

    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1))

    assert outcome.reason is Reason.INTERNAL
    sink = compose(outcome)
    assert sink.done is not None and sink.done.data["status"] == DoneStatus.ERROR.value
    assert sink.of_type(EventType.ERROR)


def test_section_timeout_is_reported_as_failed_section():
    executor = ReportFakeExecutor(delays={"get_geography": 0.3})
    orchestrator, _ = build(
        executor,
        plan=plan_json("REPORT_WEEKLY"),
        config=OrchestratorConfig(step_timeout_seconds=0.05, report_turn_timeout_seconds=30),
    )

    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    by_id = {section.section_id: section.status for section in outcome.sections}

    assert by_id["geography"] is SectionStatus.ERROR
    assert by_id["units"] is SectionStatus.OK


def test_all_sections_without_data_give_no_data():
    executor = ReportFakeExecutor(no_data=list(ReportFakeExecutor().rows))
    orchestrator, _ = build(executor, plan=plan_json("REPORT_WEEKLY"))

    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1))
    sink = compose(outcome)

    assert sink.done is not None
    assert sink.done.data["status"] == DoneStatus.NO_DATA.value


def test_custom_report_without_dates_asks_for_range():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_CUSTOM"))
    outcome = orchestrator.handle(turn("Làm báo cáo tuỳ chỉnh cho tôi", TV1))
    assert outcome.decision is Decision.CLARIFY
    assert outcome.reason is Reason.REPORT_RANGE_REQUIRED
    sink = compose(outcome)
    assert sink.of_type(EventType.CLARIFY)


def test_report_sections_filter_keeps_overview():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_CUSTOM", ["products"]))
    outcome = orchestrator.handle(turn("Báo cáo từ 01/08 đến 31/08 về sản phẩm", TV1))
    assert [section.section_id for section in outcome.sections] == ["overview", "products"]


def test_report_steps_do_not_exceed_max_steps():
    orchestrator, _ = build(
        ReportFakeExecutor(),
        plan=plan_json("REPORT_WEEKLY"),
        config=OrchestratorConfig(report_max_steps=2),
    )
    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    statuses = [section.status for section in outcome.sections]
    assert statuses[:2] == [SectionStatus.OK, SectionStatus.OK]
    assert all(status is SectionStatus.SKIPPED for status in statuses[2:])


# ── Event của báo cáo ──


def test_highlights_block_comes_first_and_sections_are_ordered():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_WEEKLY"))
    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    sink = compose(outcome)

    data_blocks = blocks(sink)
    assert data_blocks[0]["section"]["id"] == "highlights"
    indexes = [b["section"]["index"] for b in data_blocks[1:] if "section" in b]
    assert indexes == sorted(indexes)


def test_highlights_contains_delta_of_total_issues():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_WEEKLY"))
    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    sink = compose(outcome)
    first = blocks(sink)[0]
    assert "tăng 23%" in first["payload"]["items"][0]["display_text"]


def test_non_report_answer_has_no_section_field():
    from .ai_fakes import ScopeRespectingFakeExecutor
    from .test_orchestrator import PLAN_JSON
    from .test_orchestrator import build as build_normal
    from .test_orchestrator import turn as normal_turn

    orchestrator, _ = build_normal(ScopeRespectingFakeExecutor(), responses=PLAN_JSON)
    outcome = orchestrator.handle(normal_turn("Tổng quan tháng 8", TV1))
    sink = compose(outcome)
    assert all("section" not in block for block in blocks(sink))


# ── Nhận định theo phần ──


REPORT_COMMENTARY = (
    "[overview] Số vấn đề {{overview.delta.total_issues}} so với kỳ trước.\n"
    "[products] {{products.rank.1.label}} bị phản hồi nhiều nhất, "
    "chiếm {{products.rank.1.pct}} tổng số.\n"
)


def test_report_calls_llm_once_for_commentary():
    executor = ReportFakeExecutor()
    orchestrator, llm = build(
        executor,
        plan=plan_json("REPORT_WEEKLY"),
        stream_chunks=[[REPORT_COMMENTARY]],
    )
    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    sink = compose(outcome, llm=llm, commentary=True)

    synthesis_calls = [kind for kind, _ in llm.calls if kind == "chat_synthesis"]
    assert len(synthesis_calls) == 1
    texts = [e.data["text"] for e in sink.of_type(EventType.COMMENTARY)]
    assert texts[0].startswith("Số vấn đề tăng 23%")
    assert sink.of_type(EventType.COMMENTARY)[0].data["section"]["id"] == "overview"


@pytest.mark.parametrize(
    ("sentence", "reason"),
    [
        ("[products] Số vấn đề {{overview.delta.total_issues}}.", "CROSS_SECTION_FACT"),
        ("[khong_co] Có vẻ ổn.", "UNKNOWN_SECTION"),
        ("Không có tiền tố phần nào.", "UNKNOWN_SECTION"),
    ],
)
def test_bad_report_sentences_are_dropped(sentence, reason):
    from dms.chat.ai.fact_sheet import FactSheet, make_fact
    from dms.chat.guardrails.synthesis_guard import SynthesisGuard

    sheet = FactSheet(
        facts={
            "overview.delta.total_issues": make_fact(
                "overview.delta.total_issues", "tăng 23%", 23.0
            ),
            "products.rank.1.label": make_fact("products.rank.1.label", "Đèn LED Bulb"),
        }
    )
    guard = SynthesisGuard(sheet, sections=("overview", "products"))
    result = guard.check(sentence)
    assert not result.ok and result.reason is not None and result.reason.value == reason


def test_at_most_two_sentences_per_section():
    executor = ReportFakeExecutor()
    one_line = "[overview] Số vấn đề {{overview.delta.total_issues}} so với kỳ trước.\n"
    long_commentary = one_line * 4
    orchestrator, llm = build(
        executor, plan=plan_json("REPORT_WEEKLY"), stream_chunks=[[long_commentary]]
    )
    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    sink = compose(outcome, llm=llm, commentary=True)
    assert len(sink.of_type(EventType.COMMENTARY)) == 2


def test_report_survives_llm_failure():
    executor = ReportFakeExecutor()
    orchestrator, llm = build(
        executor,
        plan=plan_json("REPORT_WEEKLY"),
        stream_chunks=[[RuntimeError("gemini down")]],
    )
    outcome = orchestrator.handle(turn("Báo cáo tuần trước", TV1, TV2))
    sink = compose(outcome, llm=llm, commentary=True)

    assert sink.done is not None
    assert sink.done.data["commentary_status"] == "unavailable"
    assert blocks(sink)[0]["section"]["id"] == "highlights"


# ── Xuất Excel (b10 D7, task 4.2, 4.6) ──


def exporter(tmp_path):
    from dms.chat.ai.export.export_store import ExportStore
    from dms.chat.ai.export.exporter import ChatExporter

    return ChatExporter(ExportStore(tmp_path / "chat_exports"))


def compose_with_export(outcome, tmp_path, *, llm=None):
    sink = ListSink()
    AnswerComposer(
        llm=llm,
        config=ComposerConfig(commentary_enabled=False, milestone="M5"),
        exporter=exporter(tmp_path),
    ).compose(outcome, sink)
    return sink


def export_payload(sink: ListSink) -> dict:
    found = [b for b in blocks(sink) if b["kind"] == "export"]
    assert found, "không có khối export"
    return found[0]["payload"]


def test_export_report_produces_a_downloadable_file(tmp_path):
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_EXPORT"))
    outcome = orchestrator.handle(turn("Xuất báo cáo từ 01/09 đến 07/09 ra Excel", TV1))

    assert outcome.export_request is True
    sink = compose_with_export(outcome, tmp_path)
    payload = export_payload(sink)

    assert payload["export_id"] and "error" not in payload
    assert payload["filename"].endswith(".xlsx")
    assert payload["size_bytes"] > 0
    stored = tmp_path / "chat_exports"
    assert list(stored.rglob("*.xlsx"))


def test_export_without_previous_answer_asks_what_to_export():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_EXPORT"))
    outcome = orchestrator.handle(turn("Xuất kết quả vừa rồi ra Excel", TV1))

    assert outcome.decision is Decision.CLARIFY
    assert outcome.reason is Reason.EXPORT_NO_SOURCE
    assert outcome.message == "Bạn muốn xuất dữ liệu nào?"


def test_export_of_previous_answer_uses_the_current_scope(tmp_path):
    previous_plans = [
        {
            "step": 1,
            "function_name": "get_units",
            "params": {"date_from": "2026-09-01", "date_to": "2026-09-07", "unit_name": TV1},
            "intent": "DRILL_UNIT",
        },
        {
            "step": 2,
            "function_name": "get_units",
            "params": {"date_from": "2026-09-01", "date_to": "2026-09-07", "unit_name": TV2},
            "intent": "DRILL_UNIT",
        },
    ]
    executor = ReportFakeExecutor()
    orchestrator, _ = build(executor, plan=plan_json("REPORT_EXPORT"))

    # Quyền đã bị thu hẹp còn TV1 sau khi hỏi.
    outcome = orchestrator.handle(
        turn("Xuất kết quả vừa rồi ra Excel", TV1, previous_plans=tuple(previous_plans))
    )

    assert outcome.decision is Decision.RUN
    units = {plan.params.get("unit_name") for plan, _ in executor.calls}
    assert units == {TV1}
    assert any(TV2 in notice.message for notice in outcome.notices)

    sink = compose_with_export(outcome, tmp_path)
    assert export_payload(sink)["export_id"]
    exported = next(iter((tmp_path / "chat_exports").rglob("*.xlsx")))
    from openpyxl import load_workbook

    text = "\n".join(
        str(cell.value or "")
        for row in load_workbook(exported)["Thông tin"].iter_rows()
        for cell in row
    )
    assert TV2 in text and "không có quyền" in text


def test_export_of_previous_answer_never_reuses_saved_events(tmp_path):
    """Dữ liệu trong file luôn được lấy lại, không lấy từ events đã lưu."""
    executor = ReportFakeExecutor()
    orchestrator, _ = build(executor, plan=plan_json("REPORT_EXPORT"))
    plans = [
        {
            "step": 1,
            "function_name": "get_products",
            "params": {"date_from": "2026-09-01", "date_to": "2026-09-07"},
            "intent": "DRILL_PRODUCT",
        }
    ]
    orchestrator.handle(turn("Xuất kết quả vừa rồi ra Excel", TV1, previous_plans=tuple(plans)))
    assert [plan.function_name for plan, _ in executor.calls] == ["get_products"]


def test_export_block_reports_error_without_exporter():
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_EXPORT"))
    outcome = orchestrator.handle(turn("Xuất báo cáo từ 01/09 đến 07/09 ra Excel", TV1))
    sink = compose(outcome)
    assert export_payload(sink)["error"] == "EXPORT_FAILED"


def test_export_of_a_weekly_range_is_named_as_a_weekly_report(tmp_path):
    orchestrator, _ = build(ReportFakeExecutor(), plan=plan_json("REPORT_EXPORT"))
    outcome = orchestrator.handle(turn("Xuất báo cáo tuần trước ra Excel", TV1))
    sink = compose_with_export(outcome, tmp_path)
    assert export_payload(sink)["filename"] == "bao-cao-tuan_2026-09-07_2026-09-13.xlsx"
