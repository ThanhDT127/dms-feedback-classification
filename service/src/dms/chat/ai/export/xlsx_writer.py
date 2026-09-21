"""Ghi file ``.xlsx`` từ các khối của một câu trả lời (spec ``chat-report-export``, b10 D8).

Nguyên tắc: số là ô số, ngày là ô ngày, và **mọi ô chữ nguy hiểm đều được thêm tiền tố ``'``**
để Excel không bao giờ coi nội dung của khách hàng là công thức.
"""

from __future__ import annotations

import logging
import os
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..vi_format import NOT_ENOUGH_DATA

logger = logging.getLogger("dms-chat-export")

INFO_SHEET_TITLE = "Thông tin"
SHEET_TITLE_MAX_CHARS = 31
# Excel coi ô chữ mở đầu bằng các ký tự này là công thức.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
INT_FORMAT = "#,##0"
PERCENT_FORMAT = "0.0%"
DATE_FORMAT = "dd/mm/yyyy"
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
AI_COMMENTARY_NOTE = "Phần nhận định do trợ lý AI viết từ số liệu trong file này."


class ExportTooLarge(Exception):
    """File vượt ``CHAT_EXPORT_MAX_BYTES``; bản tạm đã bị xoá."""


@dataclass(frozen=True)
class ExportInfo:
    """Nội dung sheet "Thông tin"."""

    title: str
    exported_by: str
    exported_at: str
    scope: str
    date_range: str = ""
    filters: Sequence[str] = ()
    assumptions: Sequence[str] = ()
    notes: Sequence[str] = ()
    highlights: Sequence[str] = ()
    commentary: Sequence[str] = ()


@dataclass
class WriteResult:
    path: Path
    size_bytes: int = 0
    rows_exported: int = 0
    truncated: bool = False
    sheets: list[str] = field(default_factory=list)


def sanitize_cell(value: Any) -> Any:
    """Giá trị cho ô chữ: chặn công thức bằng tiền tố ``'`` (design D8)."""
    text = "" if value is None else str(value)
    if text.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def sheet_title(raw: str, used: set[str]) -> str:
    """Tên sheet hợp lệ, tối đa 31 ký tự, không trùng."""
    cleaned = _INVALID_SHEET_CHARS.sub(" ", str(raw or "")).strip() or "Dữ liệu"
    cleaned = " ".join(cleaned.split())[:SHEET_TITLE_MAX_CHARS]
    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used:
        tail = f" ({suffix})"
        candidate = cleaned[: SHEET_TITLE_MAX_CHARS - len(tail)].rstrip() + tail
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def write_workbook(
    path: Path,
    info: ExportInfo,
    blocks: Sequence[Mapping[str, Any]],
    *,
    max_rows: int = 5000,
    max_bytes: int = 10 * 1024 * 1024,
) -> WriteResult:
    """Ghi ra file tạm cùng thư mục rồi ``os.replace``, để không để lại file hỏng."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".part")

    workbook = Workbook()
    used: set[str] = set()
    info_sheet = workbook.active
    info_sheet.title = sheet_title(INFO_SHEET_TITLE, used)
    _write_info(info_sheet, info)

    rows_exported = 0
    truncated = False
    sheets: list[str] = []
    for block in blocks:
        written, cut = _write_block(workbook, block, used, max_rows=max_rows)
        rows_exported += written
        truncated = truncated or cut
        sheets.append(workbook.sheetnames[-1])

    if truncated:
        _append_note(info_sheet, f"Dữ liệu bị cắt ở {max_rows:,} dòng mỗi bảng.".replace(",", "."))

    try:
        workbook.save(temp)
        size = temp.stat().st_size
        if size > max_bytes:
            raise ExportTooLarge(f"{size} > {max_bytes}")
        os.replace(temp, path)
    except ExportTooLarge:
        temp.unlink(missing_ok=True)
        raise
    except Exception:
        temp.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()

    return WriteResult(
        path=path,
        size_bytes=path.stat().st_size,
        rows_exported=rows_exported,
        truncated=truncated,
        sheets=sheets,
    )


# ── Sheet "Thông tin" ──


def _write_info(sheet: Worksheet, info: ExportInfo) -> None:
    _set_text(sheet, 1, 1, info.title, bold=True)
    row = 3
    pairs = (
        ("Người xuất", info.exported_by),
        ("Thời điểm xuất", info.exported_at),
        ("Phạm vi", info.scope),
        ("Khoảng thời gian", info.date_range),
    )
    for label, value in pairs:
        _set_text(sheet, row, 1, label, bold=True)
        _set_text(sheet, row, 2, value)
        row += 1

    for label, values in (
        ("Bộ lọc", info.filters),
        ("Giả định", info.assumptions),
        ("Ghi chú", info.notes),
        ("Điểm chính", info.highlights),
    ):
        row = _write_list(sheet, row, label, values)
    if info.commentary:
        row = _write_list(sheet, row, "Nhận định tự động", info.commentary)
        _set_text(sheet, row, 2, AI_COMMENTARY_NOTE)
        row += 1

    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 90


def _write_list(sheet: Worksheet, row: int, label: str, values: Sequence[str]) -> int:
    if not values:
        return row
    row += 1
    _set_text(sheet, row, 1, label, bold=True)
    for value in values:
        _set_text(sheet, row, 2, value)
        row += 1
    return row


def _append_note(sheet: Worksheet, text: str) -> None:
    _set_text(sheet, sheet.max_row + 2, 1, "Ghi chú", bold=True)
    _set_text(sheet, sheet.max_row, 2, text)


def _set_text(sheet: Worksheet, row: int, column: int, value: Any, *, bold: bool = False) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.value = sanitize_cell(value)
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if bold:
        cell.font = Font(bold=True)


# ── Sheet dữ liệu ──


def _write_block(
    workbook: Workbook, block: Mapping[str, Any], used: set[str], *, max_rows: int
) -> tuple[int, bool]:
    kind = str(block.get("kind") or "")
    title = str(block.get("title") or "Dữ liệu")
    sheet = workbook.create_sheet(sheet_title(title, used))
    _set_text(sheet, 1, 1, title, bold=True)
    subtitle = str(block.get("subtitle") or "")
    if subtitle:
        _set_text(sheet, 2, 1, subtitle)

    payload = block.get("payload") or {}
    header_row = 4
    if kind == "kpi":
        columns, rows = _kpi_table(payload)
    elif kind == "ranking":
        columns, rows = _ranking_table(payload)
    elif kind == "timeseries":
        columns, rows = _timeseries_table(payload)
    elif kind == "quote":
        columns, rows = _quote_table(payload)
    else:
        columns, rows = _plain_table(payload)

    truncated = len(rows) > max_rows
    rows = rows[:max_rows]
    for index, (header, _) in enumerate(columns, start=1):
        _set_text(sheet, header_row, index, header, bold=True)
    for offset, record in enumerate(rows, start=header_row + 1):
        for index, (_, key) in enumerate(columns, start=1):
            _write_value(sheet, offset, index, record.get(key))
    _autosize(sheet, columns)
    return len(rows), truncated


def _write_value(sheet: Worksheet, row: int, column: int, value: Any) -> None:
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, _Number):
        cell.value = value.value
        cell.number_format = value.number_format
        return
    if isinstance(value, date | datetime):
        cell.value = value
        cell.number_format = DATE_FORMAT
        return
    if isinstance(value, str) and _ISO_DATE.match(value):
        cell.value = date.fromisoformat(value)
        cell.number_format = DATE_FORMAT
        return
    cell.value = sanitize_cell(value)
    cell.data_type = "s"


@dataclass(frozen=True)
class _Number:
    value: int | float
    number_format: str = INT_FORMAT


def _count(value: Any) -> Any:
    if isinstance(value, bool) or value is None:
        return NOT_ENOUGH_DATA
    if isinstance(value, int | float):
        return _Number(value, INT_FORMAT)
    return value


def _percent(value: Any) -> Any:
    """Excel lưu phần trăm dạng tỉ lệ: 12,5% → 0.125 với định dạng ``0.0%``."""
    if isinstance(value, bool) or value is None:
        return NOT_ENOUGH_DATA
    if isinstance(value, int | float):
        return _Number(float(value) / 100.0, PERCENT_FORMAT)
    return value


def _kpi_table(payload: Mapping[str, Any]) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    columns = [("Chỉ số", "label"), ("Giá trị", "value"), ("Kỳ trước", "previous")]
    rows = []
    for item in payload.get("items") or []:
        comparison = item.get("comparison") or {}
        rows.append(
            {
                "label": item.get("label") or item.get("key") or "",
                "value": _count(item.get("value"))
                if item.get("available")
                else str(item.get("display") or NOT_ENOUGH_DATA),
                "previous": _count(comparison.get("value"))
                if comparison.get("available")
                else "",
            }
        )
    return columns, rows


def _ranking_table(
    payload: Mapping[str, Any],
) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    columns = [
        ("#", "rank"),
        (str(payload.get("label_title") or "Tên"), "label"),
        (str(payload.get("value_title") or "Số lượng"), "value"),
        ("Tỉ lệ", "percent"),
    ]
    rows = [
        {
            "rank": _count(item.get("rank")),
            "label": item.get("label") or "",
            "value": _count(item.get("value")),
            "percent": _percent(item.get("percent")),
        }
        for item in payload.get("items") or []
    ]
    return columns, rows


def _timeseries_table(
    payload: Mapping[str, Any],
) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    columns = [("Ngày", "date"), ("Số lượng", "value")]
    rows = [
        {"date": point.get("date"), "value": _count(point.get("value"))}
        for point in payload.get("points") or []
    ]
    return columns, rows


def _quote_table(payload: Mapping[str, Any]) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    columns = [
        ("Mã vấn đề", "issue_code"),
        ("Nội dung", "content"),
        ("Đơn vị", "unit_name"),
        ("Ngày", "issue_date"),
    ]
    rows = [
        {
            "issue_code": quote.get("issue_code") or "",
            "content": quote.get("content") or "",
            "unit_name": quote.get("unit_name") or "",
            "issue_date": quote.get("issue_date"),
        }
        for quote in payload.get("quotes") or []
    ]
    return columns, rows


def _plain_table(payload: Mapping[str, Any]) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    columns_spec = payload.get("columns") or []
    columns = [
        (str(column.get("header_vi") or column.get("key")), str(column.get("key")))
        for column in columns_spec
    ]
    formats = {str(c.get("key")): str(c.get("format") or "text") for c in columns_spec}
    rows: list[dict[str, Any]] = []
    for record in payload.get("rows") or []:
        row: dict[str, Any] = {}
        for _, key in columns:
            value = record.get(key)
            # Khối table giữ chuỗi đã định dạng; số đã thành "1.234" nên để nguyên là ô chữ.
            row[key] = value if formats.get(key) != "date" else value
        rows.append(row)
    return columns, rows


def _autosize(sheet: Worksheet, columns: Sequence[tuple[str, str]]) -> None:
    for index, (header, _) in enumerate(columns, start=1):
        width = max(12, min(60, len(str(header)) + 6))
        sheet.column_dimensions[get_column_letter(index)].width = width


__all__ = [
    "AI_COMMENTARY_NOTE",
    "FORMULA_PREFIXES",
    "ExportInfo",
    "ExportTooLarge",
    "WriteResult",
    "sanitize_cell",
    "sheet_title",
    "write_workbook",
]
